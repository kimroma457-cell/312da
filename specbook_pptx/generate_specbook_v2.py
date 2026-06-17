"""
인테리어 스펙북 PPT 생성기 v2
네이버 이미지 검색 API 기반 — 검색어를 직접 입력하면 자재 카드를 생성합니다.

실행:
  python generate_specbook_v2.py

출력:
  specbook_v2.pptx
"""
import requests
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── 네이버 API ─────────────────────────────────────────────────────────────────
NAVER_CLIENT_ID     = "E0ioOPF01SJFcYuk6viO"
NAVER_CLIENT_SECRET = "jlXrOVV8Ws"

# ── 색상 ───────────────────────────────────────────────────────────────────────
C_BG      = RGBColor(0xF7, 0xF5, 0xF2)
C_DARK    = RGBColor(0x1A, 0x18, 0x16)
C_ACCENT  = RGBColor(0xC8, 0xA9, 0x7E)
C_MID     = RGBColor(0x8C, 0x7F, 0x74)
C_BODY    = RGBColor(0x4A, 0x45, 0x40)
C_LIGHT   = RGBColor(0xE8, 0xE2, 0xDC)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_DIVIDER = RGBColor(0xD8, 0xD2, 0xCC)

TIER_COLOR = {
    "최저가": RGBColor(0x6A, 0xA8, 0x4F),
    "보통":   RGBColor(0xC8, 0xA9, 0x7E),
    "최고가": RGBColor(0x8E, 0x44, 0xAD),
}

SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)

# 스펙 슬라이드 행 Y 위치 (최대 5행)
ROW_TOPS = [Inches(x) for x in [1.505, 2.325, 3.145, 3.965, 4.785]]
ROW_H    = Inches(0.70)


# ── 네이버 이미지 검색 ─────────────────────────────────────────────────────────

def naver_search(query: str, count: int = 3) -> list[dict]:
    """검색어로 이미지 URL 목록 반환 [{url, page}, ...]"""
    try:
        res = requests.get(
            "https://openapi.naver.com/v1/search/image",
            headers={
                "X-Naver-Client-Id":     NAVER_CLIENT_ID,
                "X-Naver-Client-Secret": NAVER_CLIENT_SECRET,
            },
            params={"query": query, "display": count, "filter": "large"},
            timeout=5,
        )
        res.raise_for_status()
        return [
            {"url": it["link"], "page": it.get("originallink", it["link"])}
            for it in res.json().get("items", [])
        ]
    except Exception as e:
        print(f"  [검색 오류] {query}: {e}")
        return []


def fetch_image(url: str) -> BytesIO | None:
    try:
        r = requests.get(url, timeout=7, headers={"User-Agent": "Mozilla/5.0"})
        if r.status_code == 200 and "image" in r.headers.get("Content-Type", ""):
            return BytesIO(r.content)
    except Exception:
        pass
    return None


# ── PPTX 헬퍼 ─────────────────────────────────────────────────────────────────

def _remove_slides(prs):
    for sid in list(prs.slides._sldIdLst):
        rid = sid.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        prs.part.drop_rel(rid)
        prs.slides._sldIdLst.remove(sid)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[0])


def _rect(slide, l, t, w, h, fill=None, line=None, lw=0.5):
    s = slide.shapes.add_shape(1, l, t, w, h)
    f = s.fill
    if fill:
        f.solid(); f.fore_color.rgb = fill
    else:
        f.background()
    ln = s.line
    if line:
        ln.color.rgb = line; ln.width = Pt(lw)
    else:
        ln.fill.background()
    return s


def _text(slide, txt, l, t, w, h,
          size=10, bold=False, color=C_DARK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p  = tf.paragraphs[0]
    p.alignment = align
    r  = p.add_run()
    r.text           = txt
    r.font.size      = Pt(size)
    r.font.bold      = bold
    r.font.color.rgb = color
    return tb


def _divider(slide, y):
    _rect(slide, Inches(0.5), y, Inches(9.0), Emu(9144), fill=C_DIVIDER)


# ── 슬라이드 제작 ─────────────────────────────────────────────────────────────

def make_cover(prs, project: dict):
    slide = _blank(prs)
    _rect(slide, 0, 0, Inches(0.06), SLIDE_H, fill=C_DARK)
    _rect(slide, Inches(0.06), 0, SLIDE_W - Inches(0.06), SLIDE_H, fill=C_BG)

    _text(slide, "2026",
          Inches(0.5), Inches(0.5), Inches(3), Inches(0.4),
          size=11, bold=True, color=C_MID)
    _text(slide, "INTERIOR",
          Inches(0.5), Inches(1.0), Inches(8), Inches(0.7),
          size=42, bold=True, color=C_DARK)
    _text(slide, "SPEC BOOK",
          Inches(0.5), Inches(1.65), Inches(8), Inches(0.7),
          size=42, bold=True, color=C_DARK)

    subtitle = (f"{project.get('name','')}  ·  "
                f"{project.get('date','')}  ·  "
                f"{project.get('designer','')}")
    _text(slide, subtitle,
          Inches(0.5), Inches(2.55), Inches(8), Inches(0.3),
          size=7, color=C_MID)

    _divider(slide, Inches(3.0))

    info = (f"위치: {project.get('location','')}   |   "
            f"면적: {project.get('area','')}   |   "
            f"공사기간: {project.get('period','')}")
    _text(slide, info,
          Inches(0.5), Inches(3.15), Inches(9), Inches(0.35),
          size=8, color=C_BODY)


def make_spec_slide(prs, space_name: str, space_en: str, items: list[dict]):
    """
    items: 최대 5개
    각 item: {item_code, tier, product, spec, finish, vendor, image_url, source_url}
    """
    slide = _blank(prs)
    _rect(slide, 0, 0, Inches(0.06), SLIDE_H, fill=C_DARK)
    _rect(slide, Inches(0.06), 0, SLIDE_W - Inches(0.06), SLIDE_H, fill=C_BG)

    # 섹션 라벨
    _text(slide, "SPACE SPECIFICATION",
          Inches(0.5), Inches(0.28), Inches(9), Inches(0.26),
          size=7, bold=True, color=C_MID)
    _divider(slide, Inches(0.64))

    # 공간명
    _text(slide, f"{space_name} 공간 스펙  ·  {space_en}",
          Inches(0.5), Inches(0.70), Inches(8.5), Inches(0.52),
          size=22, bold=True, color=C_DARK)

    # 컬럼 헤더
    _text(slide, "품목",       Inches(1.38), Inches(1.26), Inches(2.8), Inches(0.20), size=7, bold=True, color=C_MID)
    _text(slide, "재질 / 마감", Inches(4.45), Inches(1.26), Inches(3.5), Inches(0.20), size=7, bold=True, color=C_MID)
    _text(slide, "비고",       Inches(8.80), Inches(1.26), Inches(0.70), Inches(0.20), size=7, bold=True, color=C_MID)
    _divider(slide, Inches(1.46))

    for i, item in enumerate(items[:5]):
        rt = ROW_TOPS[i]
        tier        = item.get("tier", "보통")
        badge_color = TIER_COLOR.get(tier, C_ACCENT)

        # 이미지
        img_data = fetch_image(item.get("image_url", "")) if item.get("image_url") else None
        if img_data:
            try:
                slide.shapes.add_picture(img_data, Inches(0.5), rt, Inches(0.70), ROW_H)
            except Exception:
                img_data = None
        if not img_data:
            _rect(slide, Inches(0.5), rt, Inches(0.70), ROW_H, fill=C_LIGHT)
            _text(slide, "+", Inches(0.5), rt, Inches(0.70), ROW_H,
                  size=18, color=C_MID, align=PP_ALIGN.CENTER)

        # 등급 뱃지
        _rect(slide, Inches(1.38), rt + Inches(0.02), Inches(0.55), Inches(0.17), fill=badge_color)
        _text(slide, tier,
              Inches(1.38), rt + Inches(0.01), Inches(0.55), Inches(0.19),
              size=6, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

        # 품목 코드
        _text(slide, item.get("item_code", "").upper(),
              Inches(1.38), rt + Inches(0.21), Inches(1.5), Inches(0.18),
              size=7, bold=True, color=C_MID)

        # 제품명
        _text(slide, item.get("product", ""),
              Inches(1.38), rt + Inches(0.38), Inches(2.9), Inches(0.28),
              size=11, bold=True, color=C_DARK)

        # 규격
        _text(slide, item.get("spec", ""),
              Inches(1.38), rt + Inches(0.53), Inches(2.9), Inches(0.20),
              size=7, color=C_MID)

        # 재질 / 마감
        _text(slide, item.get("finish", ""),
              Inches(4.45), rt + Inches(0.12), Inches(3.6), Inches(0.26),
              size=9, color=C_BODY)

        # 비고
        _text(slide, item.get("vendor", ""),
              Inches(8.80), rt + Inches(0.12), Inches(0.70), Inches(0.26),
              size=8, color=C_MID)

        # 출처
        src = item.get("source_url", "")
        if src:
            _text(slide, f"출처: {src[:55]}",
                  Inches(4.45), rt + Inches(0.44), Inches(4.3), Inches(0.20),
                  size=6, color=RGBColor(0x33, 0x66, 0xCC))

        if i < len(items) - 1:
            _divider(slide, rt + ROW_H + Inches(0.015))


# ── 대화형 입력 ───────────────────────────────────────────────────────────────

def ask(prompt: str, default: str = "") -> str:
    val = input(f"{prompt} [{default}]: ").strip()
    return val if val else default


def input_items(space_name: str) -> list[dict]:
    """공간 내 자재 항목을 입력받아 네이버 검색 후 반환."""
    items = []
    print(f"\n  [{space_name}] 자재 입력 (빈 줄 입력 시 완료)")
    while True:
        print(f"\n  자재 {len(items)+1}번 (엔터만 치면 완료)")
        search_q = input("  검색어: ").strip()
        if not search_q:
            break

        item_code = ask("  품목 코드 (예: SOFA, FLOOR, LIGHT)", "ITEM")
        product   = ask("  제품명")
        spec      = ask("  규격 (예: W1200 × D600 × H420mm)")
        finish    = ask("  재질/마감 (예: 오크 원목 / 내추럴 오일)")
        vendor    = ask("  비고 (예: 기성품/커스텀/수입)")
        tier      = ask("  등급 [최저가/보통/최고가]", "보통")

        print(f"  🔍 네이버 이미지 검색 중: '{search_q}'")
        results = naver_search(search_q, count=3)

        image_url  = ""
        source_url = ""

        if results:
            print(f"  검색 결과 {len(results)}개:")
            for ri, r in enumerate(results):
                print(f"    [{ri+1}] {r['url'][:70]}")
            choice = input("  사용할 이미지 번호 (1~3, 엔터=1): ").strip()
            idx = int(choice) - 1 if choice.isdigit() else 0
            idx = max(0, min(idx, len(results) - 1))
            image_url  = results[idx]["url"]
            source_url = results[idx]["page"]
            print(f"  ✓ 이미지 선택: {image_url[:60]}")
        else:
            print("  ⚠️  이미지를 찾지 못했습니다.")

        items.append({
            "item_code":  item_code,
            "tier":       tier,
            "product":    product,
            "spec":       spec,
            "finish":     finish,
            "vendor":     vendor,
            "image_url":  image_url,
            "source_url": source_url,
        })
        print(f"  ✓ '{product}' 추가됨")

    return items


# ── 메인 ─────────────────────────────────────────────────────────────────────

def main():
    print("=" * 55)
    print("  인테리어 스펙북 PPT 생성기 v2")
    print("  네이버 이미지 검색 API 기반")
    print("=" * 55)

    # 프로젝트 정보
    print("\n[1] 프로젝트 정보")
    project = {
        "name":     ask("프로젝트명",  "○○ 아파트 리모델링"),
        "location": ask("위치",        "서울시 강남구"),
        "area":     ask("면적",        "84㎡"),
        "period":   ask("공사기간",    "2026.07 ~ 2026.08"),
        "designer": ask("담당자",      "홍길동"),
        "date":     ask("작성일",      "2026.06.17"),
    }

    # 공간 및 자재
    print("\n[2] 공간 및 자재 입력")
    spaces = []
    SPACE_MAP = {
        "1": ("거실",     "Living Room"),
        "2": ("침실",     "Bedroom"),
        "3": ("주방",     "Kitchen & Dining"),
        "4": ("욕실",     "Bathroom"),
        "5": ("현관",     "Entrance"),
        "6": ("드레스룸", "Dressing Room"),
        "7": ("서재",     "Study Room"),
    }

    while True:
        print("\n공간 추가:")
        for k, (ko, en) in SPACE_MAP.items():
            print(f"  {k}. {ko} ({en})")
        print("  0. 직접 입력   |   엔터 = 완료")
        choice = input("선택: ").strip()

        if not choice:
            break
        elif choice == "0":
            ko = input("공간명 (한글): ").strip()
            en = input("공간명 (영어): ").strip()
        elif choice in SPACE_MAP:
            ko, en = SPACE_MAP[choice]
            # 중복 카운트
            same = sum(1 for s in spaces if s["name"].startswith(ko))
            if same:
                ko = f"{ko}{same + 1}"
        else:
            print("잘못된 선택입니다.")
            continue

        items = input_items(ko)
        spaces.append({"name": ko, "name_en": en, "items": items})
        print(f"\n✓ {ko} 완료 ({len(items)}개 자재)")

    if not spaces:
        print("공간이 없어 종료합니다.")
        return

    # PPT 생성
    print("\n[3] PPT 생성 중...")
    template_path = "../specbook_app/template.pptx"
    try:
        prs = Presentation(template_path)
        _remove_slides(prs)
        print(f"  템플릿 로드: {template_path}")
    except Exception:
        prs = Presentation()
        prs.slide_width  = SLIDE_W
        prs.slide_height = SLIDE_H
        print("  템플릿 없음 — 기본 슬라이드 사용")

    make_cover(prs, project)
    print("  ✓ 표지 생성")

    for space in spaces:
        items = space["items"]
        for chunk in range(0, max(len(items), 1), 5):
            make_spec_slide(prs, space["name"], space["name_en"],
                            items[chunk:chunk + 5])
        print(f"  ✓ {space['name']} 슬라이드 생성")

    output = "specbook_v2.pptx"
    prs.save(output)
    print(f"\n✅ 완료: {output}")
    print(f"   총 {len(prs.slides)}장 슬라이드")


if __name__ == "__main__":
    main()
